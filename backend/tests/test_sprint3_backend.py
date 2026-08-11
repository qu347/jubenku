import csv
import io
import zipfile
from datetime import date
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session
from sqlalchemy.exc import SQLAlchemyError

from app.core.config import settings
from app.models import GenreMetric, GenreModule, Material, MaterialTag, Tag
from app.services import material_service as material_service_module
from app.services.genre_metric_service import IMPORT_HEADERS


def create_module(db: Session, name: str = "测试题材", slug: str = "test-genre") -> GenreModule:
    module = GenreModule(name=name, slug=slug)
    db.add(module)
    db.commit()
    return module


def upload(
    client: TestClient,
    module: GenreModule,
    files: list[tuple[str, bytes, str]] | None = None,
    **data,
):
    payload = {
        "genre_module_id": module.id,
        "material_type": "参考资料",
        "tags": '["设定", "重点"]',
        "source": "内部资料",
        "description": "测试说明",
        **data,
    }
    multipart = [
        ("files", (filename, content, mime_type))
        for filename, content, mime_type in (files or [("sample.txt", b"hello", "text/plain")])
    ]
    return client.post("/api/materials/upload", data=payload, files=multipart)


def metric_payload(module: GenreModule, **changes) -> dict:
    return {
        "genre_module_id": module.id,
        "platform": "番茄小说",
        "channel": "男频",
        "period": "2026-08",
        "average_age": 28,
        "age_group": "middle",
        "education_level": "medium",
        "audience_share": 36.5,
        "heat_index": 82,
        "trend": "rising",
        "is_core": True,
        "sample_size": 1200,
        "data_source": "内部调研",
        "remark": "测试记录",
        **changes,
    }


def csv_import_bytes(rows: list[list[object]]) -> bytes:
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(IMPORT_HEADERS)
    writer.writerows(rows)
    return output.getvalue().encode("utf-8-sig")


def valid_import_row(module_name: str, **changes) -> list[object]:
    values = {
        "题材": module_name,
        "数据平台": "番茄小说",
        "频道": "男频",
        "数据周期": "2026-08",
        "平均年龄": 27,
        "年龄层级": "中年",
        "学历层级": "中学历",
        "用户占比": 42,
        "热度指数": 88,
        "趋势": "上升",
        "是否重点题材": "是",
        "样本数量": 500,
        "数据来源": "内部调研",
        "备注": "导入测试",
        **changes,
    }
    return [values[header] for header in IMPORT_HEADERS]


def test_single_and_multiple_upload(client: TestClient, db_session: Session) -> None:
    module = create_module(db_session)
    response = upload(client, module)
    assert response.status_code == 200
    data = response.json()["data"]
    assert data["success_count"] == 1
    assert data["materials"][0]["original_filename"] == "sample.txt"
    assert data["materials"][0]["tags"] == ["设定", "重点"]

    response = upload(
        client,
        module,
        [("one.md", b"# one", "text/markdown"), ("two.csv", b"a,b\n1,2", "text/csv")],
    )
    assert response.json()["data"]["success_count"] == 2


def test_story_metadata_and_text_content_are_returned(client: TestClient, db_session: Session) -> None:
    module = create_module(db_session)
    response = upload(
        client,
        module,
        [("story.md", "# 灰塔守钟人\n钟声响起，失踪者的名字浮现在塔壁。".encode("utf-8"), "text/markdown")],
        title="灰塔守钟人",
        material_type="剧情",
        tags='["剧情:悬疑", "时代背景:架空", "角色设定:小人物"]',
        description="守钟人发现每次钟响都会抹去一段城市记忆。",
        uploaded_by="张三",
        project_owner="李制片",
    )
    assert response.status_code == 200
    created = response.json()["data"]["materials"][0]
    assert created["title"] == "灰塔守钟人"
    assert created["uploaded_by"] == "张三"
    assert created["project_owner"] == "李制片"

    detail = client.get(f"/api/materials/{created['id']}")
    assert detail.status_code == 200
    payload = detail.json()["data"]
    assert "失踪者的名字" in payload["content_text"]
    assert payload["content_truncated"] is False

    updated = client.patch(
        f"/api/materials/{created['id']}",
        json={"uploaded_by": "王编辑", "project_owner": "赵导演"},
    )
    assert updated.status_code == 200
    assert updated.json()["data"]["uploaded_by"] == "王编辑"
    assert updated.json()["data"]["project_owner"] == "赵导演"


def test_partial_upload_failure_and_no_residue(client: TestClient, db_session: Session) -> None:
    module = create_module(db_session)
    response = upload(
        client,
        module,
        [("good.txt", b"ok", "text/plain"), ("blocked.exe", b"bad", "application/octet-stream")],
    )
    data = response.json()["data"]
    assert data["success_count"] == 1
    assert data["failure_count"] == 1
    assert "不支持" in data["results"][1]["error"]
    files = [path for path in Path(settings.material_storage_path).rglob("*") if path.is_file()]
    assert len(files) == 1
    assert not any(path.suffix == ".part" for path in files)


def test_upload_size_limit_cleans_partial_file(
    client: TestClient,
    db_session: Session,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    module = create_module(db_session)
    monkeypatch.setattr(settings, "max_upload_mb", 0)
    response = upload(client, module, [("large.txt", b"x", "text/plain")])
    assert response.json()["data"]["failure_count"] == 1
    assert not any(path.is_file() for path in Path(settings.material_storage_path).rglob("*"))


def test_mime_mismatch_is_rejected(client: TestClient, db_session: Session) -> None:
    module = create_module(db_session)
    response = upload(client, module, [("image.png", b"not png", "text/plain")])
    assert response.json()["data"]["failure_count"] == 1
    assert "MIME" in response.json()["data"]["results"][0]["error"]


@pytest.mark.parametrize(
    ("filename", "mime_type", "content"),
    [
        ("spoof.pdf", "application/pdf", b"not a pdf"),
        ("spoof.png", "image/png", b"not a png"),
        ("spoof.jpg", "image/jpeg", b"not a jpeg"),
        (
            "spoof.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            b"PK\x03\x04not-a-valid-zip",
        ),
        (
            "spoof.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            b"PK\x03\x04not-a-valid-zip",
        ),
        ("binary.txt", "text/plain", b"safe-prefix\x00binary"),
    ],
)
def test_spoofed_file_signatures_are_rejected(
    client: TestClient,
    db_session: Session,
    filename: str,
    mime_type: str,
    content: bytes,
) -> None:
    module = create_module(db_session)
    response = upload(client, module, [(filename, content, mime_type)])
    data = response.json()["data"]
    assert data["success_count"] == 0 and data["failure_count"] == 1
    assert not any(path.is_file() for path in Path(settings.material_storage_path).rglob("*"))


def test_valid_file_signatures_are_accepted(client: TestClient, db_session: Session) -> None:
    module = create_module(db_session)

    def office_file(entries: dict[str, bytes]) -> bytes:
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w") as archive:
            for name, content in entries.items():
                archive.writestr(name, content)
        return output.getvalue()

    docx = office_file({"[Content_Types].xml": b"<Types/>", "word/document.xml": b"<document/>"})
    xlsx = office_file({"[Content_Types].xml": b"<Types/>", "xl/workbook.xml": b"<workbook/>"})
    response = upload(
        client,
        module,
        [
            ("valid.pdf", b"%PDF-1.7\n%%EOF", "application/pdf"),
            ("valid.png", b"\x89PNG\r\n\x1a\n", "image/png"),
            ("valid.jpg", b"\xff\xd8\xff\xe0", "image/jpeg"),
            (
                "valid.docx",
                docx,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ),
            (
                "valid.xlsx",
                xlsx,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        ],
    )
    assert response.json()["data"]["success_count"] == 5


def test_path_traversal_filename_cannot_escape_storage(client: TestClient, db_session: Session) -> None:
    module = create_module(db_session)
    response = upload(client, module, [("../../escape.txt", b"safe", "text/plain")])
    material = response.json()["data"]["materials"][0]
    assert material["original_filename"] == "escape.txt"
    assert ".." not in material["storage_path"]
    target = (Path(settings.material_storage_path) / material["storage_path"]).resolve()
    target.relative_to(Path(settings.material_storage_path).resolve())
    assert target.read_bytes() == b"safe"


def test_upload_rejects_resolved_storage_directory_escape(
    client: TestClient,
    db_session: Session,
    upload_storage_path: Path,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    module = create_module(db_session)
    storage_root = upload_storage_path.resolve()
    outside = tmp_path / "outside-upload-target"
    outside.mkdir()
    original_resolve = Path.resolve

    def redirected_resolve(path: Path, *args, **kwargs) -> Path:
        if path != storage_root and storage_root in path.parents:
            return outside / path.relative_to(storage_root)
        return original_resolve(path, *args, **kwargs)

    monkeypatch.setattr(Path, "resolve", redirected_resolve)

    response = upload(client, module, [("safe.txt", b"safe", "text/plain")])

    data = response.json()["data"]
    assert data["success_count"] == 0
    assert data["failure_count"] == 1
    assert "越界" in data["results"][0]["error"]
    assert not any(path.is_file() for path in outside.rglob("*"))


def test_material_list_pagination_filters_and_keyword(client: TestClient, db_session: Session) -> None:
    first = create_module(db_session)
    second = create_module(db_session, "另一题材", "other-genre")
    upload(client, first, [("alpha.txt", b"a", "text/plain")], source="甲来源")
    upload(client, first, [("beta.md", b"b", "text/markdown")], source="乙来源")
    upload(client, second, [("gamma.txt", b"c", "text/plain")], source="甲来源")

    page = client.get("/api/materials", params={"page": 1, "page_size": 2}).json()["data"]
    assert page["total"] == 3 and page["pages"] == 2 and len(page["items"]) == 2
    by_genre = client.get("/api/materials", params={"genre_module_id": first.id}).json()["data"]
    assert by_genre["total"] == 2
    by_type = client.get("/api/materials", params={"file_extension": "md"}).json()["data"]
    assert by_type["total"] == 1
    keyword = client.get("/api/materials", params={"keyword": "gamma"}).json()["data"]
    assert keyword["total"] == 1
    tags = client.get("/api/materials", params={"tags": "重点", "source": "甲来源"}).json()["data"]
    assert tags["total"] == 2
    today = date.today().isoformat()
    by_date = client.get(
        "/api/materials", params={"uploaded_from": today, "uploaded_to": today}
    ).json()["data"]
    assert by_date["total"] == 3


def test_material_list_requires_all_selected_tags(client: TestClient, db_session: Session) -> None:
    module = create_module(db_session)
    upload(
        client,
        module,
        [("both.md", b"both", "text/markdown")],
        tags='["剧情:逆袭", "角色设定:真假千金"]',
    )
    upload(
        client,
        module,
        [("plot-only.md", b"plot", "text/markdown")],
        tags='["剧情:逆袭", "角色设定:小人物"]',
    )
    upload(
        client,
        module,
        [("role-only.md", b"role", "text/markdown")],
        tags='["剧情:马甲", "角色设定:真假千金"]',
    )

    response = client.get(
        "/api/materials",
        params={"tags": "剧情:逆袭,角色设定:真假千金"},
    )
    assert response.status_code == 200
    data = response.json()["data"]
    assert data["total"] == 1
    assert data["items"][0]["original_filename"] == "both.md"


def test_legacy_content_material_is_visible_without_attachment(client: TestClient, db_session: Session) -> None:
    legacy = Material(
        title="旧素材",
        material_type="legacy",
        summary="旧素材中文摘要",
        content="旧素材中文正文，必须在升级后继续可见。",
    )
    legacy_tag = Tag(name="旧标签")
    db_session.add_all((legacy, legacy_tag))
    db_session.flush()
    db_session.add(MaterialTag(material_id=legacy.id, tag_id=legacy_tag.id))
    db_session.commit()
    listed = client.get("/api/materials").json()["data"]
    assert listed["total"] == 1
    assert listed["items"][0]["has_attachment"] is False
    assert listed["items"][0]["legacy_summary"] == "旧素材中文摘要"
    assert listed["items"][0]["legacy_content"] == "旧素材中文正文，必须在升级后继续可见。"
    assert listed["items"][0]["tags"] == ["旧标签"]
    assert client.get("/api/materials", params={"keyword": "中文正文"}).json()["data"]["total"] == 1
    assert client.get("/api/materials", params={"tags": "旧标签"}).json()["data"]["total"] == 1
    cleared = client.patch(f"/api/materials/{legacy.id}", json={"tags": []})
    assert cleared.status_code == 200
    assert cleared.json()["data"]["tags"] == []
    assert client.get(f"/api/materials/{legacy.id}").json()["data"]["tags"] == []
    assert client.get("/api/materials", params={"tags": "旧标签"}).json()["data"]["total"] == 0
    detail = client.get(f"/api/materials/{legacy.id}")
    assert detail.status_code == 200
    assert detail.json()["data"]["has_attachment"] is False
    download = client.get(f"/api/materials/{legacy.id}/download")
    assert download.status_code == 404
    assert download.json()["error"]["code"] == "material_file_not_found"


def test_legacy_material_delete_never_resolves_a_storage_path(
    client: TestClient,
    db_session: Session,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    legacy = Material(title="可安全删除的旧素材", material_type="legacy")
    db_session.add(legacy)
    db_session.commit()

    def forbidden_file_access(*args, **kwargs):
        raise AssertionError("无附件记录不应访问任何磁盘文件")

    monkeypatch.setattr(material_service_module.os, "replace", forbidden_file_access)
    response = client.delete(f"/api/materials/{legacy.id}")
    assert response.status_code == 200
    assert response.json()["data"]["file_deleted"] is False


def test_update_metadata_does_not_replace_file(client: TestClient, db_session: Session) -> None:
    module = create_module(db_session)
    created = upload(client, module).json()["data"]["materials"][0]
    response = client.patch(
        f"/api/materials/{created['id']}",
        json={"title": "新标题", "tags": ["新标签"], "description": "新说明"},
    )
    assert response.status_code == 200
    updated = response.json()["data"]
    assert updated["title"] == "新标题" and updated["stored_filename"] == created["stored_filename"]
    forbidden = client.patch(
        f"/api/materials/{created['id']}", json={"stored_filename": "tampered.txt"}
    )
    assert forbidden.status_code == 422


def test_download_and_physical_delete(client: TestClient, db_session: Session) -> None:
    module = create_module(db_session)
    created = upload(client, module, [("download.txt", b"download body", "text/plain")]).json()["data"]["materials"][0]
    path = Path(settings.material_storage_path) / created["storage_path"]
    response = client.get(f"/api/materials/{created['id']}/download")
    assert response.status_code == 200 and response.content == b"download body"
    assert "download.txt" in response.headers["content-disposition"]
    deleted = client.delete(f"/api/materials/{created['id']}")
    assert deleted.status_code == 200 and deleted.json()["data"]["file_deleted"] is True
    assert not path.exists()
    assert client.get(f"/api/materials/{created['id']}").status_code == 404


def test_delete_unlink_failure_restores_record_and_file(
    client: TestClient,
    db_session: Session,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    module = create_module(db_session)
    created = upload(client, module).json()["data"]["materials"][0]
    path = Path(settings.material_storage_path) / created["storage_path"]
    original_unlink = Path.unlink

    def fail_tombstone_unlink(target: Path, *args, **kwargs):
        if target.name.endswith(".deleting"):
            raise OSError("forced unlink failure")
        return original_unlink(target, *args, **kwargs)

    monkeypatch.setattr(Path, "unlink", fail_tombstone_unlink)
    response = client.delete(f"/api/materials/{created['id']}")
    assert response.status_code == 500
    assert response.json()["error"]["code"] == "material_file_cleanup_failed"
    assert path.read_bytes() == b"hello"
    assert client.get(f"/api/materials/{created['id']}").status_code == 200
    assert not any(item.name.endswith(".deleting") for item in path.parent.iterdir())

    monkeypatch.setattr(Path, "unlink", original_unlink)
    assert client.delete(f"/api/materials/{created['id']}").status_code == 200
    assert not path.exists()


def test_delete_commit_failure_restores_record_and_file(
    client: TestClient,
    db_session: Session,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    module = create_module(db_session)
    created = upload(client, module).json()["data"]["materials"][0]
    path = Path(settings.material_storage_path) / created["storage_path"]
    original_commit = db_session.commit

    def fail_commit() -> None:
        raise SQLAlchemyError("forced commit failure")

    monkeypatch.setattr(db_session, "commit", fail_commit)
    response = client.delete(f"/api/materials/{created['id']}")
    assert response.status_code == 500
    assert response.json()["error"]["code"] == "material_delete_failed"
    monkeypatch.setattr(db_session, "commit", original_commit)
    assert path.read_bytes() == b"hello"
    assert client.get(f"/api/materials/{created['id']}").status_code == 200
    assert not any(item.name.endswith(".deleting") for item in path.parent.iterdir())


def test_create_update_delete_metric(client: TestClient, db_session: Session) -> None:
    module = create_module(db_session)
    response = client.post("/api/genre-metrics", json=metric_payload(module))
    assert response.status_code == 201
    metric = response.json()["data"]
    assert metric["genre_module"]["name"] == module.name
    response = client.patch(
        f"/api/genre-metrics/{metric['id']}",
        json={"average_age": 31, "education_level": "high", "audience_share": 50},
    )
    assert response.json()["data"]["average_age"] == 31
    assert response.json()["data"]["education_level"] == "high"
    assert client.delete(f"/api/genre-metrics/{metric['id']}").status_code == 200
    assert client.get(f"/api/genre-metrics/{metric['id']}").status_code == 404


@pytest.mark.parametrize(
    ("changes", "field"),
    [
        ({"average_age": 0}, "average_age"),
        ({"average_age": 101}, "average_age"),
        ({"audience_share": -1}, "audience_share"),
        ({"audience_share": 101}, "audience_share"),
        ({"heat_index": 101}, "heat_index"),
        ({"sample_size": -1}, "sample_size"),
        ({"period": "近30天"}, "period"),
    ],
)
def test_metric_validation(
    client: TestClient,
    db_session: Session,
    changes: dict,
    field: str,
) -> None:
    module = create_module(db_session)
    response = client.post("/api/genre-metrics", json=metric_payload(module, **changes))
    assert response.status_code == 422
    assert any(field in item["field"] for item in response.json()["error"]["details"])


def test_metric_filters_and_pagination(client: TestClient, db_session: Session) -> None:
    module = create_module(db_session)
    client.post("/api/genre-metrics", json=metric_payload(module, platform="番茄小说", education_level="medium"))
    client.post("/api/genre-metrics", json=metric_payload(module, platform="起点中文网", education_level="high", heat_index=55))
    data = client.get(
        "/api/genre-metrics",
        params={"platform": "番茄小说", "education_level": "medium", "heat_min": 80},
    ).json()["data"]
    assert data["total"] == 1 and data["items"][0]["platform"] == "番茄小说"
    assert client.get(
        "/api/genre-metrics", params={"heat_min": 90, "heat_max": 10}
    ).status_code == 422


def test_legacy_metric_values_are_canonicalized_without_migration_rewrite(
    client: TestClient,
    db_session: Session,
) -> None:
    module = create_module(db_session)
    metric = GenreMetric(
        genre_module_id=module.id,
        platform="旧平台",
        channel="旧频道",
        period="近30天",
        average_age=30,
        age_group="中年",
        education_level="高学历",
        audience_share=50,
        heat_index=80,
        trend="stable",
        is_core=True,
        sample_size=100,
        data_source="旧来源",
        remark="",
    )
    db_session.add(metric)
    db_session.commit()
    response = client.get(
        "/api/genre-metrics",
        params={"age_group": "middle", "education_level": "high", "period": "2026-08"},
    )
    item = response.json()["data"]["items"][0]
    assert item["age_group"] == "middle"
    assert item["education_level"] == "high"
    assert item["period"] == "2026-08"
    db_session.refresh(metric)
    assert (metric.age_group, metric.education_level, metric.period) == ("中年", "高学历", "近30天")


def test_csv_import_success_and_partial_errors(client: TestClient, db_session: Session) -> None:
    module = create_module(db_session)
    content = csv_import_bytes(
        [
            valid_import_row(module.name),
            valid_import_row(module.name, **{"用户占比": 120}),
            valid_import_row("不存在的题材"),
        ]
    )
    response = client.post(
        "/api/genre-metrics/import", files={"file": ("metrics.csv", content, "text/csv")}
    )
    data = response.json()["data"]
    assert data["success_count"] == 1 and data["failure_count"] == 2
    assert {item["row"] for item in data["errors"]} == {3, 4}
    assert any(item["field"] == "用户占比" for item in data["errors"])
    assert any("不存在" in item["reason"] for item in data["errors"])


def test_xlsx_import_success(client: TestClient, db_session: Session) -> None:
    module = create_module(db_session)
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(IMPORT_HEADERS)
    sheet.append(valid_import_row(module.name))
    output = io.BytesIO()
    workbook.save(output)
    response = client.post(
        "/api/genre-metrics/import",
        files={
            "file": (
                "metrics.xlsx",
                output.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.json()["data"] == {"success_count": 1, "failure_count": 0, "errors": []}


def test_import_templates_and_exports(client: TestClient, db_session: Session) -> None:
    module = create_module(db_session)
    client.post("/api/genre-metrics", json=metric_payload(module))
    csv_template = client.get("/api/genre-metrics/import-template", params={"format": "csv"})
    assert csv_template.status_code == 200
    assert csv_template.content.decode("utf-8-sig").splitlines()[0].split(",") == IMPORT_HEADERS
    xlsx_template = client.get("/api/genre-metrics/import-template", params={"format": "xlsx"})
    assert list(load_workbook(io.BytesIO(xlsx_template.content)).active.values)[0] == tuple(IMPORT_HEADERS)

    exported_csv = client.get(
        "/api/genre-metrics/export", params={"format": "csv", "platform": "番茄小说"}
    )
    decoded = exported_csv.content.decode("utf-8-sig")
    assert "测试题材" in decoded and "内部调研" in decoded
    exported_xlsx = client.get("/api/genre-metrics/export", params={"format": "xlsx"})
    rows = list(load_workbook(io.BytesIO(exported_xlsx.content)).active.values)
    assert rows[0] == tuple(IMPORT_HEADERS) and rows[1][0] == module.name


def test_spreadsheet_formula_injection_is_exported_as_text(
    client: TestClient,
    db_session: Session,
) -> None:
    module = create_module(db_session, "=2+3", "formula-safe")
    response = client.post(
        "/api/genre-metrics",
        json=metric_payload(
            module,
            platform="+SUM(A1:A2)",
            channel="-danger",
            data_source="@external",
            remark='=HYPERLINK("https://example.invalid")',
        ),
    )
    assert response.status_code == 201

    csv_response = client.get("/api/genre-metrics/export", params={"format": "csv"})
    csv_rows = list(csv.reader(io.StringIO(csv_response.content.decode("utf-8-sig"))))
    assert csv_rows[1][0].startswith("'=")
    assert csv_rows[1][1].startswith("'+")
    assert csv_rows[1][2].startswith("'-")
    assert csv_rows[1][12].startswith("'@")
    assert csv_rows[1][13].startswith("'=")

    xlsx_response = client.get("/api/genre-metrics/export", params={"format": "xlsx"})
    sheet = load_workbook(io.BytesIO(xlsx_response.content), data_only=False).active
    for cell in (sheet["A2"], sheet["B2"], sheet["C2"], sheet["M2"], sheet["N2"]):
        assert str(cell.value).startswith("'")
        assert cell.data_type == "s"

    template = client.get("/api/genre-metrics/import-template", params={"format": "xlsx"})
    header_cells = next(
        load_workbook(io.BytesIO(template.content), data_only=False).active.iter_rows()
    )
    assert all(cell.data_type == "s" for cell in header_cells)
