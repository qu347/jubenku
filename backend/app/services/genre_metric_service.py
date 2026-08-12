from __future__ import annotations

import csv
import io
import math
from datetime import datetime, timezone
from typing import Any, Iterable

from fastapi import UploadFile, status
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from pydantic import ValidationError
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.exceptions import AppException
from app.models import GenreMetric, GenreModule
from app.repositories.genre_metric_repository import GenreMetricRepository
from app.schemas.genre_metric import GenreMetricCreate, GenreMetricUpdate, validate_period


IMPORT_HEADERS = [
    "题材",
    "数据平台",
    "频道",
    "数据周期",
    "平均年龄",
    "年龄层级",
    "学历层级",
    "用户占比",
    "热度指数",
    "趋势",
    "是否重点题材",
    "样本数量",
    "数据来源",
    "备注",
]

FIELD_HEADERS = {
    "genre_module_id": "题材",
    "platform": "数据平台",
    "channel": "频道",
    "period": "数据周期",
    "average_age": "平均年龄",
    "age_group": "年龄层级",
    "education_level": "学历层级",
    "audience_share": "用户占比",
    "heat_index": "热度指数",
    "trend": "趋势",
    "is_core": "是否重点题材",
    "sample_size": "样本数量",
    "data_source": "数据来源",
    "remark": "备注",
}

AGE_ALIASES = {"少年": "youth", "中年": "middle", "老年": "senior"}
EDUCATION_ALIASES = {"低学历": "low", "中学历": "medium", "高学历": "high"}
TREND_ALIASES = {"上升": "rising", "稳定": "stable", "下降": "falling"}
PERIOD_ALIASES = {"近30天": "2026-08"}
FORMULA_PREFIXES = ("=", "+", "-", "@")


def safe_spreadsheet_value(value: Any) -> Any:
    """Keep untrusted text from being interpreted as a spreadsheet formula."""
    if isinstance(value, str) and value.startswith(FORMULA_PREFIXES):
        return f"'{value}"
    return value


def metric_to_dict(metric: GenreMetric) -> dict[str, Any]:
    module = metric.genre_module
    return {
        "id": metric.id,
        "genre_module_id": metric.genre_module_id,
        "genre_module": {
            "id": module.id,
            "name": module.name,
            "slug": module.slug,
            "theme_color": module.theme_color,
        },
        "platform": metric.platform,
        "channel": metric.channel,
        "period": PERIOD_ALIASES.get(metric.period, metric.period),
        "average_age": metric.average_age,
        "age_group": AGE_ALIASES.get(metric.age_group, metric.age_group),
        "education_level": EDUCATION_ALIASES.get(metric.education_level, metric.education_level),
        "audience_share": metric.audience_share,
        "heat_index": metric.heat_index,
        "trend": TREND_ALIASES.get(metric.trend, metric.trend),
        "is_core": metric.is_core,
        "sample_size": metric.sample_size,
        "data_source": metric.data_source,
        "remark": metric.remark,
        "created_at": metric.created_at,
        "updated_at": metric.updated_at,
        "deleted_at": metric.deleted_at,
    }


class GenreMetricService:
    def __init__(self, session: Session) -> None:
        self.session = session
        self.repository = GenreMetricRepository(session)

    def _get(self, metric_id: str) -> GenreMetric:
        metric = self.repository.get(metric_id)
        if not metric:
            raise AppException(
                "定位数据不存在",
                status_code=status.HTTP_404_NOT_FOUND,
                code="genre_metric_not_found",
            )
        return metric

    def _require_module(self, module_id: str) -> GenreModule:
        module = self.repository.get_module(module_id)
        if not module:
            raise AppException(
                "题材模块不存在",
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                code="genre_module_not_found",
                details={"field": "genre_module_id"},
            )
        return module

    @staticmethod
    def _validate_filters(filters: dict[str, Any]) -> None:
        heat_min = filters.get("heat_min")
        heat_max = filters.get("heat_max")
        if heat_min is not None and heat_max is not None and heat_min > heat_max:
            raise AppException(
                "最低热度不能高于最高热度",
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                code="invalid_heat_range",
            )
        if filters.get("period"):
            try:
                filters["period"] = validate_period(filters["period"])
            except ValueError as exc:
                raise AppException(
                    str(exc),
                    status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                    code="invalid_period",
                ) from exc

    def list(self, *, page: int, page_size: int, **filters) -> dict[str, Any]:
        self._validate_filters(filters)
        items, total = self.repository.list(page=page, page_size=page_size, **filters)
        return {
            "items": [metric_to_dict(item) for item in items],
            "total": total,
            "page": page,
            "page_size": page_size,
            "pages": math.ceil(total / page_size) if total else 0,
        }

    def get(self, metric_id: str) -> dict[str, Any]:
        return metric_to_dict(self._get(metric_id))

    def create(self, payload: GenreMetricCreate) -> dict[str, Any]:
        values = payload.model_dump()
        values["genre_module_id"] = str(values["genre_module_id"])
        self._require_module(values["genre_module_id"])
        try:
            metric = self.repository.create(values)
            self.session.commit()
        except SQLAlchemyError as exc:
            self.session.rollback()
            raise AppException(
                "定位数据创建失败",
                status_code=status.HTTP_409_CONFLICT,
                code="genre_metric_create_failed",
            ) from exc
        return metric_to_dict(self._get(metric.id))

    def update(self, metric_id: str, payload: GenreMetricUpdate) -> dict[str, Any]:
        metric = self._get(metric_id)
        values = payload.model_dump(exclude_unset=True)
        if "genre_module_id" in values:
            values["genre_module_id"] = str(values["genre_module_id"])
            metric.genre_module = self._require_module(values["genre_module_id"])
        for field, value in values.items():
            setattr(metric, field, value)
        try:
            self.session.commit()
        except SQLAlchemyError as exc:
            self.session.rollback()
            raise AppException(
                "定位数据更新失败",
                status_code=status.HTTP_409_CONFLICT,
                code="genre_metric_update_failed",
            ) from exc
        return metric_to_dict(self._get(metric_id))

    def delete(self, metric_id: str) -> dict[str, str]:
        metric = self._get(metric_id)
        metric.soft_delete()
        self.session.commit()
        return {"id": metric_id}

    @staticmethod
    def _workbook_bytes(headers: list[str], rows: Iterable[list[Any]] = ()) -> bytes:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "题材定位数据"
        sheet.append([safe_spreadsheet_value(value) for value in headers])
        for row in rows:
            sheet.append([safe_spreadsheet_value(value) for value in row])
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        fill = PatternFill("solid", fgColor="1F4E78")
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = fill
        sheet.freeze_panes = "A2"
        widths = [18, 18, 14, 14, 12, 14, 14, 12, 12, 12, 16, 12, 20, 28]
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[chr(64 + index)].width = width
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()

    @staticmethod
    def _csv_bytes(headers: list[str], rows: Iterable[list[Any]] = ()) -> bytes:
        output = io.StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow([safe_spreadsheet_value(value) for value in headers])
        for row in rows:
            writer.writerow([safe_spreadsheet_value(value) for value in row])
        return output.getvalue().encode("utf-8-sig")

    def template(self, file_format: str) -> tuple[bytes, str, str]:
        if file_format == "csv":
            return self._csv_bytes(IMPORT_HEADERS), "text/csv; charset=utf-8", "genre_metric_import_template.csv"
        return (
            self._workbook_bytes(IMPORT_HEADERS),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "genre_metric_import_template.xlsx",
        )

    @staticmethod
    def _read_csv(content: bytes) -> list[tuple[int, dict[str, Any]]]:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise AppException(
                "CSV 文件必须使用 UTF-8 编码",
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                code="invalid_csv_encoding",
            ) from exc
        reader = csv.DictReader(io.StringIO(text))
        headers = reader.fieldnames or []
        GenreMetricService._validate_headers(headers)
        return [(index, dict(row)) for index, row in enumerate(reader, start=2)]

    @staticmethod
    def _read_xlsx(content: bytes) -> list[tuple[int, dict[str, Any]]]:
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sheet = workbook.active
            row_iterator = sheet.iter_rows(values_only=True)
            headers = [str(value).strip() if value is not None else "" for value in next(row_iterator, ())]
            GenreMetricService._validate_headers(headers)
            rows = [
                (row_number, dict(zip(headers, values)))
                for row_number, values in enumerate(row_iterator, start=2)
            ]
            workbook.close()
            return rows
        except AppException:
            raise
        except Exception as exc:
            raise AppException(
                "Excel 文件无法读取或已经损坏",
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                code="invalid_excel_file",
            ) from exc

    @staticmethod
    def _validate_headers(headers: list[str]) -> None:
        missing = [header for header in IMPORT_HEADERS if header not in headers]
        if missing:
            raise AppException(
                "导入文件缺少必需列",
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                code="missing_import_columns",
                details={"missing": missing},
            )

    @staticmethod
    def _parse_bool(value: Any) -> bool:
        if isinstance(value, bool):
            return value
        normalized = str(value or "").strip().casefold()
        if normalized in {"是", "true", "1", "yes", "y"}:
            return True
        if normalized in {"否", "false", "0", "no", "n", ""}:
            return False
        raise ValueError("只能填写是或否")

    @staticmethod
    def _validation_reason(error: dict[str, Any]) -> str:
        error_type = str(error.get("type") or "")
        if error_type == "missing":
            return "不能为空"
        if error_type in {"float_parsing", "float_type"}:
            return "必须是数字"
        if error_type in {"int_parsing", "int_type"}:
            return "必须是整数"
        if error_type == "greater_than":
            return "必须大于 0"
        if error_type in {"less_than_equal", "greater_than_equal"}:
            return "数值超出允许范围"
        if error_type == "literal_error":
            return "值不在允许范围内"
        message = str(error.get("msg") or "输入值不合法")
        if "数据周期" in message:
            return "数据周期格式应为 YYYY-MM 或 YYYY-Q1"
        return "输入值不合法"

    def _row_payload(self, row: dict[str, Any]) -> GenreMetricCreate:
        module_name = str(row.get("题材") or "").strip()
        module = self.repository.get_module_by_name(module_name) if module_name else None
        if not module:
            raise AppException(
                "题材名称不存在，不会自动创建",
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                code="import_genre_not_found",
                details={"field": "题材"},
            )
        age_group = str(row.get("年龄层级") or "").strip().lower()
        education = str(row.get("学历层级") or "").strip().lower()
        trend = str(row.get("趋势") or "").strip().lower()
        try:
            is_core = self._parse_bool(row.get("是否重点题材"))
        except ValueError as exc:
            raise AppException(
                str(exc),
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                code="invalid_import_boolean",
                details={"field": "是否重点题材"},
            ) from exc
        values = {
            "genre_module_id": module.id,
            "platform": row.get("数据平台"),
            "channel": row.get("频道"),
            "period": str(row.get("数据周期") or "").strip(),
            "average_age": row.get("平均年龄"),
            "age_group": AGE_ALIASES.get(age_group, age_group),
            "education_level": EDUCATION_ALIASES.get(education, education),
            "audience_share": row.get("用户占比"),
            "heat_index": row.get("热度指数"),
            "trend": TREND_ALIASES.get(trend, trend),
            "is_core": is_core,
            "sample_size": row.get("样本数量") if row.get("样本数量") not in {None, ""} else 0,
            "data_source": str(row.get("数据来源") or "").strip(),
            "remark": str(row.get("备注") or "").strip(),
        }
        return GenreMetricCreate.model_validate(values)

    async def import_file(self, upload: UploadFile) -> dict[str, Any]:
        filename = (upload.filename or "").replace("\\", "/").rsplit("/", 1)[-1]
        extension = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
        if extension not in {".xlsx", ".csv"}:
            raise AppException(
                "只支持 .xlsx 和 .csv 导入文件",
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                code="unsupported_import_format",
            )
        content = await upload.read(settings.max_upload_bytes + 1)
        if len(content) > settings.max_upload_bytes:
            raise AppException(
                f"导入文件超过 {settings.max_upload_mb}MB 大小限制",
                status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                code="import_file_too_large",
            )
        rows = self._read_csv(content) if extension == ".csv" else self._read_xlsx(content)
        success_count = 0
        failure_rows: set[int] = set()
        errors: list[dict[str, Any]] = []
        for row_number, row in rows:
            if all(value is None or str(value).strip() == "" for value in row.values()):
                continue
            try:
                payload = self._row_payload(row)
                values = payload.model_dump()
                values["genre_module_id"] = str(values["genre_module_id"])
                self.repository.create(values)
                self.session.commit()
                success_count += 1
            except AppException as exc:
                self.session.rollback()
                failure_rows.add(row_number)
                field = (exc.details or {}).get("field", "题材") if isinstance(exc.details, dict) else "题材"
                errors.append({"row": row_number, "field": field, "reason": exc.message})
            except ValidationError as exc:
                self.session.rollback()
                failure_rows.add(row_number)
                for item in exc.errors():
                    field_name = str(item.get("loc", ["数据"])[0])
                    errors.append(
                        {
                            "row": row_number,
                            "field": FIELD_HEADERS.get(field_name, field_name),
                            "reason": self._validation_reason(item),
                        }
                    )
            except (ValueError, SQLAlchemyError) as exc:
                self.session.rollback()
                failure_rows.add(row_number)
                reason = str(exc) if isinstance(exc, ValueError) else "数据库写入失败"
                errors.append({"row": row_number, "field": "数据", "reason": reason})
        return {
            "success_count": success_count,
            "failure_count": len(failure_rows),
            "errors": errors,
        }

    def export(self, file_format: str, **filters) -> tuple[bytes, str, str]:
        self._validate_filters(filters)
        metrics = self.repository.export(**filters)
        age_labels = {"youth": "少年", "middle": "中年", "senior": "老年"}
        education_labels = {"low": "低学历", "medium": "中学历", "high": "高学历"}
        trend_labels = {"rising": "上升", "stable": "稳定", "falling": "下降"}
        rows = [
            [
                metric.genre_module.name,
                metric.platform,
                metric.channel,
                PERIOD_ALIASES.get(metric.period, metric.period),
                metric.average_age,
                age_labels.get(metric.age_group, metric.age_group),
                education_labels.get(metric.education_level, metric.education_level),
                metric.audience_share,
                metric.heat_index,
                trend_labels.get(metric.trend, metric.trend),
                "是" if metric.is_core else "否",
                metric.sample_size,
                metric.data_source,
                metric.remark,
            ]
            for metric in metrics
        ]
        stamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
        if file_format == "csv":
            return (
                self._csv_bytes(IMPORT_HEADERS, rows),
                "text/csv; charset=utf-8",
                f"genre_metrics_{stamp}.csv",
            )
        return (
            self._workbook_bytes(IMPORT_HEADERS, rows),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            f"genre_metrics_{stamp}.xlsx",
        )
